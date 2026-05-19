import os
import re
import csv
import json
import time
import base64
import shutil
import random
import string
import mimetypes
import traceback
import subprocess
from collections import Counter
from pathlib import Path
from datetime import datetime, timedelta
from typing import Dict, List, Any, Generator, Optional, Tuple
from urllib.parse import urlparse, parse_qs

import numpy as np
import requests
import pandas as pd
import openpyxl
from openpyxl import load_workbook, Workbook
from PyPDF2 import PdfMerger
from PIL import Image

import duckdb
from rapidfuzz.distance import JaroWinkler

import pythoncom
try:
    import win32com.client
    HAS_WIN32COM = True
except ImportError:
    HAS_WIN32COM = False

try:
    from docx2pdf import convert as docx2pdf_convert
    HAS_DOCX2PDF = True
except Exception:
    HAS_DOCX2PDF = False
    
try:
    from unidecode import unidecode
    HAS_UNIDECODE = True
except ImportError:
    HAS_UNIDECODE = False

try:
    import cv2
    HAS_CV2 = True
except ImportError:
    HAS_CV2 = False

try:
    import pypdfium2 as pdfium
    HAS_PDFIUM = True
except ImportError:
    HAS_PDFIUM = False

try:
    from doctr.models import ocr_predictor
    HAS_DOCTR = True
except ImportError:
    HAS_DOCTR = False

try:
    import torch
    HAS_TORCH = True
except ImportError:
    HAS_TORCH = False


LOOKUP_CHUNK_SIZE = 100_000
INPUT_CHUNK_SIZE = 200_000
PREVIEW_DEFAULT_ROWS = 1000

IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff", ".webp"}
WORD_EXTS  = {".doc", ".docx"}
PDF_EXTS   = {".pdf"}
TEMP_DIR_NAME = "._merge_tmp"
SUPPORTED_EXTENSIONS = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp", ".pdf"}

# Allow massive image sizes to prevent decompression bomb warning
Image.MAX_IMAGE_PIXELS = None

# =========================================================
# STRUCTURED LOGGING & HELPERS
# =========================================================
def now_str() -> str: return datetime.now().strftime("%H:%M:%S")

def push_log(log_queue: List[Dict], level: str, message: str) -> None:
    log_queue.append({"timestamp": now_str(), "level": level.upper(), "message": message})

def ensure_dir(path: str) -> None:
    if path: os.makedirs(path, exist_ok=True)

def safe_file_exists(path: str) -> bool:
    return bool(path) and os.path.exists(path)

def get_file_ext(path: str) -> str:
    return os.path.splitext(path)[1].lower()


# =========================================================
# FILE READING HELPERS
# =========================================================
def smart_read(file_path: str) -> pd.DataFrame:
    ext = os.path.splitext(file_path)[1].lower()
    try:
        if ext == ".xlsx": return pd.read_excel(file_path, engine="openpyxl", dtype=str)
        elif ext == ".csv": return pd.read_csv(file_path, dtype=str)
    except: pass
    try: return pd.read_csv(file_path, dtype=str)
    except: pass
    raise ValueError("Unsupported or corrupted file")

def read_csv_preview(path: str, nrows: int) -> pd.DataFrame:
    return pd.read_csv(path, dtype=str, nrows=nrows, encoding="utf-8", on_bad_lines="skip", low_memory=False).fillna("")

def read_excel_preview(path: str, nrows: int, sheet_name: Optional[str] = None) -> pd.DataFrame:
    return pd.read_excel(path, dtype=str, nrows=nrows, engine="openpyxl", sheet_name=sheet_name if sheet_name else 0).fillna("")

def get_headers_from_file(path: str, sheet_name: Optional[str] = None) -> List[str]:
    ext = get_file_ext(path)
    if ext == ".csv": df = read_csv_preview(path, 5)
    elif ext in [".xlsx", ".xls"]: df = read_excel_preview(path, 5, sheet_name=sheet_name)
    else: raise ValueError(f"Unsupported file format: {path}")
    return [str(col) for col in df.columns]

def preview_file(path: str, nrows: int = PREVIEW_DEFAULT_ROWS, sheet_name: Optional[str] = None) -> pd.DataFrame:
    ext = get_file_ext(path)
    if ext == ".csv": return read_csv_preview(path, nrows)
    elif ext in [".xlsx", ".xls"]: return read_excel_preview(path, nrows, sheet_name=sheet_name)
    raise ValueError(f"Unsupported file format: {path}")

def iter_file_chunks(path: str, chunksize: int) -> Generator[pd.DataFrame, None, None]:
    ext = get_file_ext(path)
    if ext == ".csv": 
        for chunk in pd.read_csv(path, dtype=str, chunksize=chunksize, encoding="utf-8", on_bad_lines="skip", low_memory=False):
            yield chunk.fillna("")
    elif ext in [".xlsx", ".xls"]:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        headers = [str(cell.value) if cell.value else f"Col_{i}" for i, cell in enumerate(ws[1])]
        chunk = []
        for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=1):
            chunk.append(row)
            if i % chunksize == 0:
                yield pd.DataFrame(chunk, columns=headers).fillna("").astype(str)
                chunk = []
        if chunk:
            yield pd.DataFrame(chunk, columns=headers).fillna("").astype(str)
        wb.close()
    else: 
        raise ValueError(f"Unsupported input chunking format for pipeline: {path}")

def list_lookup_files(lookup_folders: List[str]) -> List[str]:
    files = []
    for folder in lookup_folders:
        if not safe_file_exists(folder) or not os.path.isdir(folder): continue
        for name in sorted(os.listdir(folder)):
            full_path = os.path.join(folder, name)
            if os.path.isfile(full_path) and get_file_ext(full_path) in [".csv", ".xlsx", ".xls"]:
                files.append(full_path)
    return files

def dedupe_preserve_order(items: List[str]) -> List[str]:
    seen = set()
    result = []
    for item in items:
        if item not in seen:
            seen.add(item)
            result.append(item)
    return result

class CsvChunkWriter:
    def __init__(self, file_path: str, headers: List[str]):
        self.file_path = file_path
        self.headers = headers
        self.first_write = True
        with open(self.file_path, "w", encoding="utf-8-sig") as f:
            pass

    def write_df(self, df: pd.DataFrame) -> None:
        if df.empty:
            if self.first_write:
                pd.DataFrame(columns=self.headers).to_csv(self.file_path, index=False, encoding="utf-8-sig")
                self.first_write = False
            return
        
        out_dict = {}
        for col in self.headers:
            if col in df.columns:
                out_dict[col] = df[col]
            else:
                out_dict[col] = ""
                
        out_df = pd.DataFrame(out_dict)
        out_df.to_csv(self.file_path, mode='a', index=False, header=self.first_write, encoding="utf-8-sig")
        self.first_write = False


# =========================================================
# MAPPING CONFIG
# =========================================================
def get_target_output_column(mapping: Dict[str, Any]) -> str:
    if mapping["destination_type"] == "existing": return mapping["target_existing_column"]
    return mapping["new_column_name"]

def get_all_dynamic_output_columns(config: Dict[str, Any]) -> List[str]:
    cols = []
    for mapping in config["main_lookup_mappings"]: cols.append(get_target_output_column(mapping))
    if "special_lookup_mappings" in config and config["special_lookup_mappings"]:
        for mapping in config["special_lookup_mappings"]: cols.append(get_target_output_column(mapping))
    return dedupe_preserve_order([c for c in cols if str(c).strip() != ""])

def build_output_headers(config: Dict[str, Any]) -> Tuple[List[str], List[str], List[str]]:
    selected_input_output_cols = config["selected_input_output_cols"]
    dynamic_cols = get_all_dynamic_output_columns(config)
    matched_headers = dedupe_preserve_order(selected_input_output_cols + dynamic_cols)
    unmatched_headers = dedupe_preserve_order(selected_input_output_cols + dynamic_cols + ["reason"])
    consolidated_headers = dedupe_preserve_order(selected_input_output_cols + dynamic_cols + ["match_status", "reason"])
    return matched_headers, unmatched_headers, consolidated_headers


# =========================================================
# FAST LOOKUP ENGINE (DUCKDB BULK JOINS)
# =========================================================

def normalize_series(series: pd.Series, key_type: str = "text") -> pd.Series:
    s = series.fillna("").astype(str).str.strip()
    if key_type == "mobile":
        s = s.str.replace(r"\D", "", regex=True).str[-10:]
        s = s.where(s.str.len() >= 10, "")
    else:
        s = s.str.lower()
    return s

def init_duckdb(db_path: str):
    conn = duckdb.connect(db_path)
    conn.execute("PRAGMA threads=%d" % max(1, os.cpu_count() or 4))
    conn.execute("PRAGMA memory_limit='8GB'")
    return conn

def build_lookup_index(config, log_queue, progress_state):
    conn = init_duckdb(config["db_path"])
    dynamic_cols = get_all_dynamic_output_columns(config)

    conn.execute("DROP TABLE IF EXISTS lookup_stage")

    cols_sql = ", ".join([f'"{c}" VARCHAR' for c in dynamic_cols])
    if not dynamic_cols:
        conn.execute("CREATE TABLE lookup_stage (norm_key VARCHAR, source_type VARCHAR, priority INTEGER)")
    else:
        conn.execute(f"CREATE TABLE lookup_stage (norm_key VARCHAR, {cols_sql}, source_type VARCHAR, priority INTEGER)")

    key_type = config.get("key_type", "text")
    sources = []

    if config.get("special_lookup_file"):
        sources.append((config["special_lookup_file"], config.get("special_lookup_key_col"), config.get("special_lookup_mappings", []), 1, "special_lookup"))

    for f in list_lookup_files(config["lookup_folders"]):
        sources.append((f, config["main_lookup_key_col"], config["main_lookup_mappings"], 2, "main_lookup"))

    total = len(sources)

    for i, (file_path, key_col, mappings, priority, src_type) in enumerate(sources, 1):
        push_log(log_queue, "INFO", f"Indexing {file_path}")

        try:
            df = smart_read(file_path).fillna("")
        except PermissionError:
            push_log(log_queue, "ERROR", f"Permission Denied. Please close {os.path.basename(file_path)} if it is open in Excel.")
            raise

        if key_col not in df.columns:
            push_log(log_queue, "WARN", f"Skipped {os.path.basename(file_path)}: Missing key column {key_col}")
            continue

        df["norm_key"] = normalize_series(df[key_col], key_type)
        df = df[df["norm_key"] != ""]

        for m in mappings:
            if m["mandatory"]:
                if m["lookup_column"] in df.columns:
                    df = df[df[m["lookup_column"]].astype(str).str.strip() != ""]

        rename_map = {"norm_key": "norm_key"}
        for m in mappings:
            rename_map[m["lookup_column"]] = get_target_output_column(m)

        df = df.rename(columns=rename_map)

        for col in dynamic_cols:
            if col not in df.columns:
                df[col] = ""

        df["source_type"] = src_type
        df["priority"] = priority

        target_cols = ["norm_key"] + dynamic_cols + ["source_type", "priority"]
        conn.register("temp_df", df[target_cols])
        conn.execute("INSERT INTO lookup_stage SELECT * FROM temp_df")
        conn.unregister("temp_df")

        progress_state["lookup_percent"] = int(i / max(total, 1) * 100)

    conn.execute("DROP TABLE IF EXISTS lookup_index")

    conn.execute(f"""
        CREATE TABLE lookup_index AS
        SELECT *
        FROM (
            SELECT *,
                   ROW_NUMBER() OVER (PARTITION BY norm_key ORDER BY priority ASC) rn
            FROM lookup_stage
        ) WHERE rn = 1
    """)

    total_keys = conn.execute("SELECT COUNT(*) FROM lookup_index").fetchone()[0]
    conn.close()

    push_log(log_queue, "SUCCESS", f"Indexed {total_keys:,} unique keys")
    return {"unique_lookup_keys_indexed": total_keys, "file_level_summary_rows": []}

def process_input_file(config, log_queue, progress_state):
    conn = init_duckdb(config["db_path"])

    input_file = config["input_file"]
    key_col = config["input_key_col"]
    key_type = config.get("key_type", "text")

    matched_headers, unmatched_headers, consolidated_headers = build_output_headers(config)

    matched_writer = CsvChunkWriter(os.path.join(config["output_folder"], "matched_output.csv"), matched_headers)
    unmatched_writer = CsvChunkWriter(os.path.join(config["output_folder"], "unmatched_output.csv"), unmatched_headers)
    consolidated_writer = CsvChunkWriter(os.path.join(config["output_folder"], "consolidated_output.csv"), consolidated_headers)

    dynamic_cols = get_all_dynamic_output_columns(config)

    stats = {"processed_rows": 0, "matched_rows": 0, "unmatched_rows": 0}
    start = time.time()

    try:
        for i, chunk in enumerate(iter_file_chunks(input_file, INPUT_CHUNK_SIZE), 1):
            chunk = chunk.fillna("")

            if key_col not in chunk.columns:
                raise ValueError(f"Input key column '{key_col}' not found in input file.")

            chunk["norm_key"] = normalize_series(chunk[key_col], key_type)

            conn.register("input_chunk", chunk)

            lookup_cols_sql = ", ".join([f'l."{c}"' for c in dynamic_cols])
            if lookup_cols_sql:
                lookup_cols_sql = lookup_cols_sql + ","

            sql = f"""
                SELECT
                    i.*,
                    {lookup_cols_sql}
                    CASE
                        WHEN i.norm_key = '' THEN 'INVALID_OR_BLANK_KEY'
                        WHEN l.norm_key IS NULL THEN 'NO_MATCH_FOUND'
                        ELSE ''
                    END AS reason
                FROM input_chunk i
                LEFT JOIN lookup_index l
                ON i.norm_key = l.norm_key
            """

            df = conn.execute(sql).df()
            conn.unregister("input_chunk")

            stats["processed_rows"] += len(df)

            matched = df[df["reason"] == ""]
            unmatched = df[df["reason"] != ""]

            stats["matched_rows"] += len(matched)
            stats["unmatched_rows"] += len(unmatched)

            matched_writer.write_df(matched)
            unmatched_writer.write_df(unmatched)

            df["match_status"] = df["reason"].apply(lambda x: "MATCHED" if x == "" else "UNMATCHED")
            consolidated_writer.write_df(df)

            speed = int(stats["processed_rows"] / max(time.time() - start, 1))

            progress_state["input_rows_processed"] = stats["processed_rows"]
            progress_state["throughput_rps"] = speed
            progress_state["input_percent"] = min(100, i * 10)

            push_log(log_queue, "INFO", f"Chunk {i} done | speed {speed:,}/s")
            
    except PermissionError:
        push_log(log_queue, "ERROR", f"Permission Denied. Please close {os.path.basename(input_file)} if it is open in Excel.")
        raise
    finally:
        conn.close()

    push_log(log_queue, "SUCCESS", "Matching completed")
    return stats

def run_engine(config: Dict[str, Any], log_queue: List[Dict], progress_state: Dict[str, Any]) -> Dict[str, Any]:
    progress_state.update({"phase": "INITIALIZATION", "lookup_percent": 0, "input_percent": 0, "lookup_rows_scanned": 0, "lookup_files_done": 0, "lookup_files_total": 0, "input_chunk_count": 0, "input_rows_processed": 0, "status": "RUNNING", "error": ""})

    try:
        push_log(log_queue, "INFO", "Validation completed. Initializing Vectorized DuckDB Engine.")
        progress_state["phase"] = "INDEXING_LOOKUP"
        lookup_stats = build_lookup_index(config, log_queue, progress_state)
        
        progress_state["phase"] = "MATCHING_BATCHES"
        input_stats = process_input_file(config, log_queue, progress_state)

        progress_state["phase"] = "WRITING_REPORTS"
        txt_path = os.path.join(config["output_folder"], "summary.txt")
        with open(txt_path, "w") as f: f.write(f"Processed: {input_stats['processed_rows']} | Matched: {input_stats['matched_rows']}")

        progress_state["phase"] = "FINALIZED"
        progress_state["status"] = "COMPLETED"
        progress_state["lookup_percent"] = 100
        progress_state["input_percent"] = 100
        
        push_log(log_queue, "SUCCESS", "Pipeline executed securely.")
        return {
            "lookup_stats": lookup_stats, "input_stats": input_stats,
            "file_level_summary_path": "N/A (Vectorized Update)", "summary_txt_path": txt_path,
            "matched_output_path": os.path.join(config["output_folder"], "matched_output.csv"),
            "unmatched_output_path": os.path.join(config["output_folder"], "unmatched_output.csv"),
            "consolidated_output_path": os.path.join(config["output_folder"], "consolidated_output.csv"),
            "db_path": config["db_path"]
        }

    except Exception as e:
        progress_state["status"] = "FAILED"
        progress_state["error"] = str(e)
        push_log(log_queue, "ERROR", f"CRITICAL CRASH: {str(e)}")
        push_log(log_queue, "ERROR", traceback.format_exc())
        raise


# =========================================================
# STANDALONE TOOLS
# =========================================================
def find_soffice_exe() -> Optional[str]:
    candidates = ["soffice"]
    if os.name == "nt":
        candidates += [
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        ]
    for c in candidates:
        try:
            result = subprocess.run([c, "--version"], stdout=subprocess.PIPE, stderr=subprocess.PIPE)
            if result.returncode == 0:
                return c
        except Exception:
            pass
    return None

def batch_convert_words(word_files: List[Path], tmp: Path, log_queue: List[Dict]) -> List[Path]:
    out = []
    soffice = find_soffice_exe()
    for doc in word_files:
        try:
            pdf_path = tmp / f"{doc.stem}.pdf"
            if pdf_path.exists():
                out.append(pdf_path)
                continue

            if HAS_DOCX2PDF and doc.suffix.lower() in {".docx", ".doc"}:
                before = set(tmp.glob("*.pdf"))
                docx2pdf_convert(str(doc), str(tmp))
                after = set(tmp.glob("*.pdf"))
                new_files = list(after - before)
                
                chosen = None
                for nf in new_files:
                    if nf.stem.lower() == doc.stem.lower():
                        chosen = nf
                        break
                if not chosen and len(new_files) == 1:
                    chosen = new_files[0]
                if chosen and chosen != pdf_path:
                    chosen.rename(pdf_path)
                if pdf_path.exists():
                    out.append(pdf_path)
                    continue
            
            if soffice:
                cmd = [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(tmp), str(doc)]
                res = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
                if res.returncode == 0:
                    produced = tmp / f"{doc.stem}.pdf"
                    if produced.exists():
                        if produced != pdf_path:
                            produced.rename(pdf_path)
                        out.append(pdf_path)
        except Exception as e:
            push_log(log_queue, "WARN", f"Word convert failed: {doc} -> {e}")
    return out

def run_pdf_merger(target_folder: str, log_queue: List[Dict]) -> bool:
    def natural_key(s: str) -> List[Any]:
        return [int(t) if t.isdigit() else t.lower() for t in re.split(r'(\d+)', s)]
        
    push_log(log_queue, "INFO", f"Starting Recursive PDF Merge in: {target_folder}")
    path = Path(target_folder).resolve()
    
    if not path.exists() or not path.is_dir():
        push_log(log_queue, "ERROR", "Invalid folder path provided.")
        return False

    folders = [path] + [p for p in path.rglob("*") if p.is_dir() and p.name != TEMP_DIR_NAME]

    for f in folders:
        push_log(log_queue, "INFO", f"--- Processing Folder: {f.name} ---")
        output_pdf = f.parent / f"{f.name}.pdf"
        tmp = f / TEMP_DIR_NAME
        tmp.mkdir(exist_ok=True)
        
        try:
            pdfs, images, words = [], [], []
            for root, dirs, files in os.walk(f):
                dirs[:] = [d for d in dirs if d != TEMP_DIR_NAME]
                for fn in files:
                    p = Path(root) / fn
                    if fn.startswith("~$") or p.name == output_pdf.name: continue
                    if p.suffix.lower() in PDF_EXTS: pdfs.append(p)
                    elif p.suffix.lower() in IMAGE_EXTS: images.append(p)
                    elif p.suffix.lower() in WORD_EXTS: words.append(p)
            
            if not (pdfs or images or words):
                push_log(log_queue, "WARN", f"No convertible files in: {f}")
                continue

            converted_images = []
            for img in images:
                try:
                    with Image.open(img) as im:
                        if im.mode in ("RGBA", "P"): im = im.convert("RGB")
                        pdf_path = tmp / f"{img.stem}.pdf"
                        im.save(pdf_path, "PDF", resolution=300.0)
                        converted_images.append(pdf_path)
                except Exception as e:
                    push_log(log_queue, "WARN", f"Image convert failed: {img.name} -> {e}")

            converted_words = batch_convert_words(words, tmp, log_queue) if words else []
            all_pdfs = sorted(pdfs + converted_images + converted_words, key=lambda p: natural_key(str(p.relative_to(f).parent)) + natural_key(p.name))
            
            merger = PdfMerger(strict=False)
            for pdf in all_pdfs: merger.append(str(pdf))
            if len(merger.pages) > 0:
                merger.write(str(output_pdf))
                push_log(log_queue, "SUCCESS", f"Created Merged PDF: {output_pdf.name}")
            merger.close()

        except Exception as e:
            push_log(log_queue, "ERROR", f"Failed on {f} -> {e}")
        finally:
            if tmp.exists(): shutil.rmtree(tmp, ignore_errors=True)

    push_log(log_queue, "SUCCESS", "PDF Merging completed across all folders.")
    return True

def run_excel_csv_merger(folder_path: str, output_file: str, log_queue: List[Dict]) -> bool:
    push_log(log_queue, "INFO", f"Scanning for Excel/CSV files in: {folder_path}")
    files = [f for f in os.listdir(folder_path) if f.lower().endswith((".xlsx", ".xls", ".csv"))]

    if len(files) < 2:
        push_log(log_queue, "ERROR", "Need at least 2 files to merge. Operation aborted.")
        return False

    dataframes = []
    for file in files:
        file_path = os.path.join(folder_path, file)
        try:
            df = smart_read(file_path)
            df["source_file"] = file
            dataframes.append(df)
            push_log(log_queue, "INFO", f"Loaded: {file} | Rows: {len(df)} | Cols: {len(df.columns)}")
        except Exception as e:
            push_log(log_queue, "ERROR", f"Skipped {file}: {e}")

    if not dataframes:
        push_log(log_queue, "ERROR", "No valid files could be loaded.")
        return False

    push_log(log_queue, "INFO", "Merging dataframes...")
    merged_df = pd.concat(dataframes, ignore_index=True)

    output_path = os.path.join(folder_path, output_file)
    
    try:
        if output_file.lower().endswith(".xlsx"):
            merged_df.to_excel(output_path, index=False)
        else:
            merged_df.to_csv(output_path, index=False, encoding="utf-8-sig")
            
        push_log(log_queue, "SUCCESS", f"Merged file saved: {output_path}")
        push_log(log_queue, "INFO", f"Total rows: {len(merged_df)} | Total columns: {len(merged_df.columns)}")
        return True
        
    except Exception as e:
        push_log(log_queue, "ERROR", f"Failed to save output file: {e}")
        return False

def run_excel_search(folder_path: str, search_terms: List[str], log_queue: List[Dict]) -> bool:
    def normalize_search_digits(value):
        if value is None: return ""
        return re.sub(r"\D", "", str(value))
        
    exact_set, digit_map = set(), {}
    for term in search_terms:
        raw = str(term).strip()
        if not raw: continue
        exact_set.add(raw)
        digits = normalize_search_digits(raw)
        if digits: digit_map[digits] = raw

    excel_files = [os.path.join(root, f) for root, _, files in os.walk(folder_path) for f in files if f.lower().endswith((".xlsx", ".xlsm", ".xltx", ".xltm"))]

    if not excel_files:
        push_log(log_queue, "ERROR", "No supported Excel files found in the target folder.")
        return False

    push_log(log_queue, "INFO", f"Total Excel files found: {len(excel_files)}")
    push_log(log_queue, "INFO", f"Searching for {len(exact_set)} distinct terms...")

    output_wb = Workbook()
    output_ws = output_wb.active
    output_ws.title = "Matched_Records"
    output_ws.append(["Matched_Term", "Source_File", "Sheet_Name", "Excel_Row_Number", "Row_Data"])

    total_matches, total_files_processed = 0, 0

    for file_path in excel_files:
        total_files_processed += 1
        push_log(log_queue, "INFO", f"[{total_files_processed}/{len(excel_files)}] Processing: {os.path.basename(file_path)}")

        try:
            wb = load_workbook(file_path, read_only=True, data_only=True)
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                    row_values = list(row)
                    row_matches = set()

                    for cell_value in row_values:
                        if cell_value is None: continue
                        raw_text = str(cell_value).strip()
                        if not raw_text: continue
                        
                        for term in exact_set:
                            if term == raw_text or term in raw_text: row_matches.add(term)
                        
                        digits_in_cell = normalize_search_digits(raw_text)
                        if digits_in_cell:
                            for search_digits, original_term in digit_map.items():
                                if search_digits in digits_in_cell: row_matches.add(original_term)

                    if row_matches:
                        row_text = " | ".join("" if v is None else str(v) for v in row_values)
                        for matched_term in sorted(row_matches):
                            output_ws.append([matched_term, file_path, sheet_name, row_idx, row_text])
                            total_matches += 1
            wb.close()
        except Exception as e:
            push_log(log_queue, "ERROR", f"Failed reading {os.path.basename(file_path)}: {e}")

    output_file = os.path.join(folder_path, f"Search_Results_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
    output_wb.save(output_file)

    push_log(log_queue, "SUCCESS", f"Search completed. Output saved to: {output_file}")
    push_log(log_queue, "INFO", f"Matches copied: {total_matches}")
    return True

def run_name_match(excel_path: str, col1: str, col2: str, out_col: Optional[str], log_queue: List[Dict]) -> bool:
    def simple_soundex(s: str) -> str:
        s = str(s).upper()
        cleaned = "".join(ch for ch in s if 'A' <= ch <= 'Z')
        if not cleaned: return ""
        result = cleaned[0]
        prev_code = ""
        for i in range(1, len(cleaned)):
            ch = cleaned[i]
            if ch in "BFPV": code = "1"
            elif ch in "CGJKQSXZ": code = "2"
            elif ch in "DT": code = "3"
            elif ch == "L": code = "4"
            elif ch in "MN": code = "5"
            elif ch == "R": code = "6"
            else: code = ""
            if code and code != prev_code: result += code
            prev_code = code
            if len(result) >= 4: break
        return result.ljust(4, "0")

    def clean_name_local(input_str: Any) -> str:
        if pd.isna(input_str): return ""
        input_str = str(input_str).strip()
        if not input_str: return ""
        for punc in [".", ",", ":", ";", "/", "\\", "-", "_", "\t"]: input_str = input_str.replace(punc, " ")
        input_str = " ".join(input_str.split())
        changed = True
        while changed:
            changed = False
            input_str = input_str.strip()
            for prefix in NAME_PREFIXES:
                p_upper = prefix.upper()
                if input_str.upper().startswith(p_upper):
                    input_str = input_str[len(p_upper):].strip()
                    changed = True
                    break
        tokens = [t for t in input_str.split() if t.strip()]
        if tokens:
            changed = True
            while changed and tokens:
                changed = False
                for word in NAME_STOP_WORDS:
                    if tokens[-1].upper() == word.upper():
                        tokens.pop()
                        changed = True
                        break
        return " ".join(tokens).strip()

    def generate_match_notes(s1: str, s2: str, sim: float) -> str:
        if not s1 and not s2: return "Both blank"
        if not s1 or not s2: return "One name missing"
        w1, w2 = s1.split(), s2.split()
        first1, first2 = w1[0], w2[0]
        last1, last2 = w1[-1], w2[-1]
        
        note = ""
        if sim >= 98: note = "Exact match"
        elif sim >= 90:
            if first1 == first2 and last1 == last2: note = "Very strong match (minor spelling differences)"
            else: note = "Very strong match"
        elif sim >= 75:
            if first1 == first2 and last1 != last2: note = "Strong match (first name identical, surname differs)"
            elif first1 == first2: note = "Strong match (first name matched)"
            else: note = "Strong partial match"
        elif sim >= 50:
            if first1 == first2: note = "Partial match (first name same)"
            else: note = "Partial / phonetic match"
        else: note = "Low similarity"

        if len(w1) != len(w2):
            if first1 == first2 or last1 == last2: note += " (possible extra or missing middle name)"
        phonetic_hint = ""
        if sim < 90:
            if simple_soundex(first1) == simple_soundex(first2): phonetic_hint = " (first name sounds similar)"
            elif simple_soundex(last1) == simple_soundex(last2): phonetic_hint = " (surname sounds similar)"
        return note + phonetic_hint

    try:
        push_log(log_queue, "INFO", f"Loading Data file: {excel_path}")
        try:
            df = smart_read(excel_path)
        except PermissionError:
            push_log(log_queue, "ERROR", f"Permission Denied. Please close {os.path.basename(excel_path)} if it is open in Excel.")
            return False

        if col1 not in df.columns or col2 not in df.columns:
            push_log(log_queue, "ERROR", f"Columns ('{col1}', '{col2}') not found.")
            return False

        push_log(log_queue, "INFO", f"Extracting columns: '{col1}' and '{col2}'")

        similarity_scores, notes = [], []
        rows_processed = 0
        push_log(log_queue, "INFO", "Calculating similarity row by row...")
        
        for idx, row in df.iterrows():
            name_1, name_2 = clean_name_local(row[col1]).upper(), clean_name_local(row[col2]).upper()
            sim_score = round(JaroWinkler.normalized_similarity(name_1, name_2) * 100.0, 1) if name_1 and name_2 else 0.0
            similarity_scores.append(sim_score if (name_1 and name_2) else "")
            notes.append(generate_match_notes(name_1, name_2, sim_score))
            
            rows_processed += 1
            if rows_processed % 5000 == 0: push_log(log_queue, "INFO", f"Processed {rows_processed} rows...")

        target_sim_col = out_col if out_col else "Similarity_Score"
        df[target_sim_col] = similarity_scores
        df["Notes"] = notes

        output_path = excel_path.replace(".xlsx", "_with_similarity.xlsx").replace(".csv", "_with_similarity.csv")
        if output_path.endswith(".csv"): df.to_csv(output_path, index=False, encoding="utf-8-sig")
        else: df.to_excel(output_path, index=False)

        push_log(log_queue, "SUCCESS", f"Row-by-row similarity and notes written. Saved: {output_path}")
        return True

    except Exception as e:
        push_log(log_queue, "ERROR", f"Unexpected Crash: {str(e)}\n\n{traceback.format_exc()}")
        return False

def run_file_zipper(folder_path: str, max_size_mb: float, log_queue: List[Dict]) -> bool:
    import zipfile
    push_log(log_queue, "INFO", f"Scanning folder for batch zipping: {folder_path}")
    
    files = [f for f in os.listdir(folder_path) if os.path.isfile(os.path.join(folder_path, f))]
    if not files:
        push_log(log_queue, "ERROR", "No files found in the specified directory.")
        return False

    batch, batch_size, zip_index, total_zipped = [], 0, 1, 0

    for file in files:
        file_path = os.path.join(folder_path, file)
        file_size = os.path.getsize(file_path) / (1024 * 1024)

        if file_size > max_size_mb:
            push_log(log_queue, "WARN", f"Skipping '{file}' (size {file_size:.2f} MB exceeds {max_size_mb} MB limit)")
            continue

        if batch_size + file_size > max_size_mb:
            zip_name = os.path.join(folder_path, f"archive_part_{zip_index}.zip")
            push_log(log_queue, "INFO", f"Packing {len(batch)} files into {os.path.basename(zip_name)}...")
            with zipfile.ZipFile(zip_name, 'w', zipfile.ZIP_DEFLATED) as zipf:
                for f in batch: zipf.write(os.path.join(folder_path, f), arcname=f)
            
            push_log(log_queue, "SUCCESS", f"Created: {os.path.basename(zip_name)}")
            zip_index += 1
            total_zipped += len(batch)
            batch, batch_size = [], 0

        batch.append(file)
        batch_size += file_size

    if batch:
        zip_name = os.path.join(folder_path, f"archive_part_{zip_index}.zip")
        push_log(log_queue, "INFO", f"Packing final {len(batch)} files into {os.path.basename(zip_name)}...")
        with zipfile.ZipFile(zip_name, 'w', zipfile.ZIP_DEFLATED) as zipf:
            for f in batch: zipf.write(os.path.join(folder_path, f), arcname=f)
        push_log(log_queue, "SUCCESS", f"Created: {os.path.basename(zip_name)}")
        total_zipped += len(batch)
        
    push_log(log_queue, "SUCCESS", f"Zip process complete. Compressed {total_zipped} files.")
    return True

def run_file_downloader(excel_path: str, url_col: str, rename_col: str, output_dir: str, file_source_type: str, log_queue: List[Dict]) -> bool:
    def sanitize_filename(name: str) -> str:
        name = re.sub(r'[<>:"/\\|?*]+', "_", str(name).strip())
        return re.sub(r"\s+", " ", name).strip(" .")

    def get_extension_from_response(response, url):
        content_type = response.headers.get("Content-Type", "").split(";")[0].strip()
        ext = mimetypes.guess_extension(content_type) if content_type else None
        if ext: return ext
        _, url_ext = os.path.splitext(urlparse(url).path)
        return url_ext if url_ext else ".bin"

    def extract_drive_file_id(url: str):
        for pattern in [r"/file/d/([a-zA-Z0-9_-]+)", r"id=([a-zA-Z0-9_-]+)", r"/d/([a-zA-Z0-9_-]+)"]:
            match = re.search(pattern, url)
            if match: return match.group(1)
        qs = parse_qs(urlparse(url).query)
        return qs["id"][0] if "id" in qs else None

    try:
        ensure_dir(output_dir)
        push_log(log_queue, "INFO", f"Reading Data File: {excel_path}")
        
        try:
            df = smart_read(excel_path)
        except PermissionError:
            push_log(log_queue, "ERROR", f"Permission Denied. Please close {os.path.basename(excel_path)} if it is open in Excel.")
            return False
        
        if url_col not in df.columns or rename_col not in df.columns:
            push_log(log_queue, "ERROR", "Selected columns not found in file.")
            return False
            
        results, timeout, headers = [], 60, {"User-Agent": "Mozilla/5.0"}
        push_log(log_queue, "INFO", f"Starting bulk download of {len(df)} files to: {output_dir}")
        
        for index, row in df.iterrows():
            url = str(row[url_col]).strip() if pd.notna(row[url_col]) else ""
            desired_name = str(row[rename_col]).strip() if pd.notna(row[rename_col]) else f"file_{index+1}"
            
            if not url:
                results.append({"row": index + 2, "url": url, "rename_value": desired_name, "status": "skipped - empty url", "saved_path": ""})
                continue
                
            try:
                clean_name_dl = sanitize_filename(desired_name)
                out_path_base = os.path.join(output_dir, clean_name_dl)
                is_drive = file_source_type == "Google Drive" or (file_source_type == "Auto-Detect" and ("drive.google.com" in url or "docs.google.com" in url))
                
                if is_drive:
                    file_id = extract_drive_file_id(url)
                    if not file_id: raise ValueError("Could not extract Google Drive file ID")
                    
                    session = requests.Session()
                    download_url = f"https://drive.google.com/uc?export=download&id={file_id}"
                    resp = session.get(download_url, stream=True, timeout=timeout, headers=headers)
                    resp.raise_for_status()
                    
                    token = next((v for k, v in resp.cookies.items() if k.startswith("download_warning")), None)
                    if token:
                        resp = session.get(download_url, params={"confirm": token}, stream=True, timeout=timeout, headers=headers)
                        resp.raise_for_status()
                        
                    if "text/html" in resp.headers.get("Content-Type", "").lower():
                        if any(x in resp.text[:500].lower() for x in ["google drive", "access denied", "sign in"]):
                            raise PermissionError("Access Denied. File may be private.")
                            
                    final_path = out_path_base + get_extension_from_response(resp, url)
                    with open(final_path, "wb") as f:
                        for chunk in resp.iter_content(chunk_size=8192):
                            if chunk: f.write(chunk)
                else:
                    with requests.get(url, stream=True, timeout=timeout, headers=headers) as resp:
                        resp.raise_for_status()
                        final_path = out_path_base + get_extension_from_response(resp, url)
                        with open(final_path, "wb") as f:
                            for chunk in resp.iter_content(chunk_size=8192):
                                if chunk: f.write(chunk)
                
                results.append({"row": index + 2, "url": url, "rename_value": desired_name, "status": "success", "saved_path": final_path})
                if (index + 1) % 10 == 0: push_log(log_queue, "INFO", f"Downloaded {index + 1} files...")
                
            except Exception as e:
                results.append({"row": index + 2, "url": url, "rename_value": desired_name, "status": f"failed - {str(e)}", "saved_path": ""})
                push_log(log_queue, "WARN", f"Failed Row {index+2}: {str(e)}")

        result_df = pd.DataFrame(results)
        result_file = os.path.join(output_dir, "download_report.xlsx")
        result_df.to_excel(result_file, index=False)
        
        push_log(log_queue, "SUCCESS", f"Download job finished. Report saved to: {result_file}")
        return True

    except Exception as e:
        push_log(log_queue, "ERROR", f"Download Job Crashed: {str(e)}\n{traceback.format_exc()}")
        return False

def run_mail_merge_tool(excel_path: str, word_path: str, output_folder: str, start_dt: datetime, log_queue: List[Dict]) -> bool:
    def clean_file_name_mm(name: str) -> str:
        if name is None: return ""
        return re.sub(r'[\\/:*?"<>|]', '_', str(name).strip())

    def get_field_value_by_name(data_source, field_name: str) -> str:
        try: return str(data_source.DataFields(field_name).Value).strip()
        except Exception: return ""

    def parse_excel_date(value):
        if pd.isna(value) or str(value).strip() == "": return None
        if isinstance(value, datetime): return value
        text = str(value).strip()
        for fmt in ["%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%d %m %Y", "%d-%b-%Y", "%d %b %Y", "%d-%B-%Y", "%d %B %Y", "%Y-%m-%d", "%Y/%m/%d", "%m/%d/%Y", "%m-%d-%Y"]:
            try: return datetime.strptime(text, fmt)
            except ValueError: continue
        return None

    if not HAS_WIN32COM:
        push_log(log_queue, "ERROR", "Missing pywin32 library. Install using: pip install pywin32")
        return False

    wb, word, main_doc = None, None, None

    try:
        push_log(log_queue, "INFO", f"Validating Excel file: {excel_path}")
        if not os.path.exists(excel_path): raise FileNotFoundError(f"Excel file not found: {excel_path}")

        try:
            wb = openpyxl.load_workbook(excel_path)
        except PermissionError:
            raise PermissionError(f"Permission Denied. Please close {os.path.basename(excel_path)} if it is open in Excel.")

        ws = wb.active
        ws.title = "Sheet1"
        sheet_name = ws.title
        max_row, max_col = ws.max_row, ws.max_column

        if max_row < 2: raise Exception("Excel file has no data rows.")

        headers = {str(ws.cell(row=1, column=c).value).strip(): c for c in range(1, max_col + 1) if ws.cell(row=1, column=c).value}
        
        uid_col = headers.get("Unique ID")
        if not uid_col:
            uid_col = max_col + 1
            ws.cell(row=1, column=uid_col).value = "Unique ID"
            max_col += 1

        init_col = headers.get("Initiation Date")
        if not init_col:
            init_col = max_col + 1
            ws.cell(row=1, column=init_col).value = "Initiation Date"
            max_col += 1

        dob_col = next((c for h, c in headers.items() if h.upper() in ["DOB", "DATE OF BIRTH"]), None)

        used_ids = {str(ws.cell(row=r, column=uid_col).value).strip() for r in range(2, max_row + 1) if ws.cell(row=r, column=uid_col).value}

        for idx, r in enumerate(range(2, max_row + 1)):
            u_cell = ws.cell(row=r, column=uid_col)
            if not u_cell.value or str(u_cell.value).strip() == "":
                new_id = ''.join(random.choices(string.digits + string.ascii_uppercase, k=25))
                while new_id in used_ids: new_id = ''.join(random.choices(string.digits + string.ascii_uppercase, k=25))
                u_cell.value = new_id
                used_ids.add(new_id)

            i_cell = ws.cell(row=r, column=init_col)
            i_cell.value = (start_dt + timedelta(minutes=idx)).strftime("%d-%m-%Y %H:%M")
            i_cell.number_format = "@"

            if dob_col:
                d_cell = ws.cell(row=r, column=dob_col)
                p_dob = parse_excel_date(d_cell.value)
                if p_dob:
                    d_cell.value = p_dob.strftime("%d %b %Y")
                    d_cell.number_format = "@"

        push_log(log_queue, "INFO", "Saving modified tracking data back to Excel...")
        try: 
            wb.save(excel_path)
        except PermissionError: 
            raise PermissionError(f"Cannot save Excel file. Please close {os.path.basename(excel_path)} if it is open in another program.")
        wb.close()
        wb = None
        push_log(log_queue, "SUCCESS", "Excel updated successfully. Booting Word COM Object...")

        ensure_dir(output_folder)
        excel_abs, word_abs, sheet_ref = os.path.abspath(excel_path), os.path.abspath(word_path), f"{sheet_name}$"
        excel_row_count = max_row - 1

        push_log(log_queue, "INFO", f"Connecting data source: {excel_abs}")
        
        try:
            pythoncom.CoInitialize()
            word = win32com.client.DispatchEx("Word.Application")
            word.Visible, word.DisplayAlerts = False, 0
        except Exception as e:
            raise Exception(f"Failed to boot Microsoft Word in background. Is Word installed? Error: {str(e)}")

        push_log(log_queue, "INFO", f"Opening Word Template: {word_abs}")
        main_doc = word.Documents.Open(word_abs, ReadOnly=False)
        main_doc.MailMerge.OpenDataSource(
            Name=excel_abs, ConfirmConversions=False, ReadOnly=True, LinkToSource=True, AddToRecentFiles=False, Revert=False, Format=0,
            Connection=f"Provider=Microsoft.ACE.OLEDB.12.0;Data Source={excel_abs};Extended Properties='Excel 12.0 Xml;HDR=YES;IMEX=1';",
            SQLStatement=f"SELECT * FROM [{sheet_ref}]"
        )

        ds, mm = main_doc.MailMerge.DataSource, main_doc.MailMerge
        for i in range(1, excel_row_count + 1):
            ds.ActiveRecord = ds.FirstRecord = ds.LastRecord = i
            file_id = get_field_value_by_name(ds, "Unique ID") or f"Record_{i:05d}"
            safe_id = clean_file_name_mm(file_id)
            pdf_path, word_out_path = os.path.abspath(os.path.join(output_folder, f"{safe_id}.pdf")), os.path.abspath(os.path.join(output_folder, f"{safe_id}.docx"))

            push_log(log_queue, "INFO", f"Merging record {i}/{excel_row_count} with Unique ID: {file_id}")
            before_docs = word.Documents.Count
            mm.Destination, mm.SuppressBlankLines = 0, True
            mm.Execute(False)

            if word.Documents.Count <= before_docs: raise Exception(f"Mail merge did not create output document for record {i}")
            
            merged_doc = word.ActiveDocument
            merged_doc.SaveAs2(word_out_path, FileFormat=12)
            time.sleep(0.3)
            merged_doc.ExportAsFixedFormat(OutputFileName=pdf_path, ExportFormat=17)
            merged_doc.Close(False)

            if not os.path.exists(pdf_path): raise Exception(f"PDF was not created for record {i}: {pdf_path}")
            if i % 10 == 0 or i == excel_row_count: push_log(log_queue, "INFO", f"Merged {i}/{excel_row_count} records")

        push_log(log_queue, "SUCCESS", f"Exported {excel_row_count} PDF and DOCX files to {output_folder}")
        return True

    except Exception as e:
        push_log(log_queue, "ERROR", f"Mail Merge Job Crashed: {str(e)}\n{traceback.format_exc()}")
        return False
    finally:
        if main_doc:
            try: main_doc.Close(False)
            except: pass
        if word:
            try: word.Quit()
            except: pass
        try: pythoncom.CoUninitialize()
        except: pass
        if wb:
            try: wb.close()
            except: pass


# =========================================================
# ADVANCED DATA PROFILER
# =========================================================
def run_advanced_profiler(file_path: str, profile_col: str, dedupe_col: str, log_queue: List[Dict], progress_state: Dict[str, Any]) -> bool:
    try:
        push_log(log_queue, "INFO", f"Starting Data Profiler for: {file_path}")
        if not os.path.exists(file_path):
            push_log(log_queue, "ERROR", f"File not found on system: {file_path}")
            return False
            
        file_size_mb = os.path.getsize(file_path) / (1024 * 1024)
        push_log(log_queue, "INFO", f"File size: {file_size_mb:.2f} MB. Initializing streaming engine...")
        
        start_time = time.time()
        
        total_count = 0
        found_count = 0
        not_found_count = 0
        value_counter = Counter()
        seen_ids = set()
        
        try:
            for chunk_idx, chunk in enumerate(iter_file_chunks(file_path, 200_000), 1):
                total_count += len(chunk)
                
                if profile_col in chunk.columns:
                    vals = chunk[profile_col].fillna("BLANK").astype(str).str.strip()
                    vals = vals.replace("", "BLANK")
                    value_counter.update(vals.tolist())
                    
                    low_vals = vals.str.lower()
                    found_count += (low_vals == "found").sum()
                    not_found_count += (low_vals == "not found").sum()
                    
                if dedupe_col == "-- Entire Row (All Columns) --":
                    row_hashes = pd.util.hash_pandas_object(chunk, index=False)
                    seen_ids.update(row_hashes.tolist())
                elif dedupe_col != "-- Skip Duplicate Check --" and dedupe_col in chunk.columns:
                    ids = chunk[dedupe_col].fillna("").astype(str).str.strip()
                    seen_ids.update(ids.tolist())
                    
                rps = int(total_count / max(time.time() - start_time, 1))
                progress_state["throughput_rps"] = rps
                progress_state["input_rows_processed"] = total_count
                
                if chunk_idx % 2 == 0:
                    push_log(log_queue, "INFO", f"Processed Chunk {chunk_idx} | Scanned: {total_count:,} | RPS: {rps:,}")

        except PermissionError:
            push_log(log_queue, "ERROR", "Permission Denied. The file is currently open in another program (like Excel). Please close it and try again.")
            return False
            
        elapsed = time.time() - start_time
        
        progress_state["final_total"] = int(total_count)
        progress_state["final_found"] = int(found_count)
        progress_state["final_not_found"] = int(not_found_count)
        progress_state["value_counts"] = dict(value_counter)
        
        if dedupe_col != "-- Skip Duplicate Check --":
            unique_count = len(seen_ids)
            progress_state["unique_count"] = unique_count
            progress_state["duplicate_count"] = total_count - unique_count
        else:
            progress_state["unique_count"] = 0
            progress_state["duplicate_count"] = 0
        
        push_log(log_queue, "SUCCESS", f"Profiling complete! Scanned {total_count:,} rows in {elapsed:.2f} seconds.")
        return True
        
    except Exception as e:
        push_log(log_queue, "ERROR", f"Profiler Crashed: {str(e)}\n{traceback.format_exc()}")
        return False


# =========================================================
# DATA CLEANING ENGINE HELPERS
# =========================================================
def dataframe_summary(df: pd.DataFrame) -> Dict[str, Any]:
    return {
        "row_count": int(len(df)),
        "column_count": int(df.shape[1]),
        "duplicate_rows": int(df.duplicated().sum()),
        "total_nulls": int(df.isna().sum().sum()),
        "memory_mb": round(df.memory_usage(deep=True).sum() / (1024 * 1024), 2),
    }

def apply_header_renames(df: pd.DataFrame, rename_map: Dict[str, str]) -> pd.DataFrame:
    valid_map = {old: new for old, new in rename_map.items() if old in df.columns and str(new).strip() != "" and old != new}
    if valid_map: df = df.rename(columns=valid_map)
    return df

def apply_value_replacements(df: pd.DataFrame, replacement_rules: List[Dict[str, Any]]) -> pd.DataFrame:
    for rule in replacement_rules:
        col = rule.get("column")
        old_val = rule.get("old_value")
        new_val = rule.get("new_value")
        match_type = rule.get("match_type", "exact")
        case_sensitive = rule.get("case_sensitive", False)

        if col not in df.columns: continue

        if match_type == "exact":
            df[col] = df[col].replace(old_val, new_val)
        elif match_type == "contains":
            series = df[col].astype(str)
            if case_sensitive: mask = series.str.contains(str(old_val), na=False, regex=False)
            else: mask = series.str.lower().str.contains(str(old_val).lower(), na=False, regex=False)
            df.loc[mask, col] = new_val
    return df

def drop_columns(df: pd.DataFrame, cols_to_drop: List[str]) -> pd.DataFrame:
    valid_cols = [c for c in cols_to_drop if c in df.columns]
    if valid_cols: df = df.drop(columns=valid_cols)
    return df

def add_columns_with_logic(df: pd.DataFrame, add_rules: List[Dict[str, Any]]) -> pd.DataFrame:
    for rule in add_rules:
        new_col = rule.get("new_column", "").strip()
        logic_type = rule.get("logic_type")

        if not new_col: continue

        try:
            if logic_type == "constant":
                df[new_col] = rule.get("value")
            elif logic_type == "from_existing_column":
                source_col = rule.get("source_column")
                if source_col in df.columns: df[new_col] = df[source_col]
            elif logic_type == "concat":
                source_cols = [c for c in rule.get("source_columns", []) if c in df.columns]
                separator = rule.get("separator", "")
                if source_cols: df[new_col] = df[source_cols].astype(str).fillna("").agg(separator.join, axis=1)
            elif logic_type == "math":
                left_col = rule.get("left_column")
                right_col = rule.get("right_column")
                operation = rule.get("operation", "add")

                if left_col in df.columns and right_col in df.columns:
                    left = pd.to_numeric(df[left_col], errors="coerce")
                    right = pd.to_numeric(df[right_col], errors="coerce")

                    if operation == "add": df[new_col] = left + right
                    elif operation == "subtract": df[new_col] = left - right
                    elif operation == "multiply": df[new_col] = left * right
                    elif operation == "divide": df[new_col] = np.where(right != 0, left / right, np.nan)
            elif logic_type == "conditional_if_else":
                source_col = rule.get("source_column")
                operator = rule.get("operator")
                compare_value = rule.get("compare_value")
                true_value = rule.get("true_value")
                false_value = rule.get("false_value")

                if source_col in df.columns:
                    s = df[source_col]
                    if operator == "==": mask = s.astype(str) == str(compare_value)
                    elif operator == "!=": mask = s.astype(str) != str(compare_value)
                    elif operator == ">": mask = pd.to_numeric(s, errors="coerce") > pd.to_numeric(compare_value, errors="coerce")
                    elif operator == "<": mask = pd.to_numeric(s, errors="coerce") < pd.to_numeric(compare_value, errors="coerce")
                    elif operator == ">=": mask = pd.to_numeric(s, errors="coerce") >= pd.to_numeric(compare_value, errors="coerce")
                    elif operator == "<=": mask = pd.to_numeric(s, errors="coerce") <= pd.to_numeric(compare_value, errors="coerce")
                    elif operator == "contains": mask = s.astype(str).str.contains(str(compare_value), case=False, na=False, regex=False)
                    elif operator == "is_null": mask = s.isna() | (s == "")
                    elif operator == "not_null": mask = s.notna() & (s != "")
                    else: mask = pd.Series(False, index=df.index)

                    df[new_col] = np.where(mask, true_value, false_value)
        except Exception:
            pass
    return df

def apply_row_filters(df: pd.DataFrame, filter_rules: List[Dict[str, Any]]) -> pd.DataFrame:
    for rule in filter_rules:
        col = rule.get("column")
        operator = rule.get("operator")
        val = rule.get("value")

        if col not in df.columns:
            continue

        s = df[col]
        try:
            if operator == "==": mask = s.astype(str) == str(val)
            elif operator == "!=": mask = s.astype(str) != str(val)
            elif operator == ">": mask = pd.to_numeric(s, errors="coerce") > pd.to_numeric(val, errors="coerce")
            elif operator == "<": mask = pd.to_numeric(s, errors="coerce") < pd.to_numeric(val, errors="coerce")
            elif operator == ">=": mask = pd.to_numeric(s, errors="coerce") >= pd.to_numeric(val, errors="coerce")
            elif operator == "<=": mask = pd.to_numeric(s, errors="coerce") <= pd.to_numeric(val, errors="coerce")
            elif operator == "contains": mask = s.astype(str).str.contains(str(val), case=False, na=False, regex=False)
            elif operator == "is_null": mask = s.isna() | (s == "")
            elif operator == "not_null": mask = s.notna() & (s != "")
            else: mask = pd.Series(True, index=df.index)

            df = df[mask]
        except Exception:
            pass
    return df

def select_final_columns(df: pd.DataFrame, final_columns: List[str]) -> pd.DataFrame:
    valid_cols = [c for c in final_columns if c in df.columns]
    if valid_cols:
        return df[valid_cols]
    return df

def apply_all_transformations(
    df: pd.DataFrame, rename_map: Dict[str, str], replacement_rules: List[Dict[str, Any]], 
    cols_to_drop: List[str], add_rules: List[Dict[str, Any]], filter_rules: List[Dict[str, Any]] = None, final_columns: List[str] = None
) -> pd.DataFrame:
    out = df.copy()
    out = apply_header_renames(out, rename_map)
    out = apply_value_replacements(out, replacement_rules)
    out = add_columns_with_logic(out, add_rules)
    out = drop_columns(out, cols_to_drop)
    
    if filter_rules:
        out = apply_row_filters(out, filter_rules)
    
    if final_columns:
        out = select_final_columns(out, final_columns)
        
    return out

def run_data_cleaner(config: Dict[str, Any], log_queue: List[Dict], progress_state: Dict[str, Any]) -> bool:
    input_file = config["input_file"]
    output_folder = config["output_folder"]
    ensure_dir(output_folder)
    
    out_file = os.path.join(output_folder, "cleaned_output.csv")
    
    rename_map = config.get("rename_map", {})
    replacement_rules = config.get("replacement_rules", [])
    cols_to_drop = config.get("cols_to_drop", [])
    add_rules = config.get("add_rules", [])
    filter_rules = config.get("filter_rules", [])
    final_columns = config.get("final_columns", [])

    file_size = os.path.getsize(input_file)
    bytes_processed_est = 0
    stats = {"processed_rows": 0, "retained_rows": 0}
    
    start_time = time.time()
    writer = None

    try:
        push_log(log_queue, "INFO", f"Starting streaming batch cleanup for: {input_file}")
        
        try:
            for chunk_idx, chunk in enumerate(iter_file_chunks(input_file, INPUT_CHUNK_SIZE), 1):
                chunk = chunk.fillna("")
                
                cleaned_chunk = apply_all_transformations(
                    chunk, rename_map, replacement_rules, cols_to_drop, add_rules, filter_rules, final_columns
                )
                
                if writer is None:
                    headers = list(cleaned_chunk.columns)
                    writer = CsvChunkWriter(out_file, headers)
                
                writer.write_df(cleaned_chunk)
                
                stats["processed_rows"] += len(chunk)
                stats["retained_rows"] += len(cleaned_chunk)
                
                bytes_processed_est += len(chunk) * 150 
                elapsed = time.time() - start_time
                rps = int(stats["processed_rows"] / max(elapsed, 1))
                remaining_rows = max(0, int((file_size / max(bytes_processed_est, 1)) * stats["processed_rows"]) - stats["processed_rows"])
                
                progress_state["input_percent"] = min(100, int((bytes_processed_est / max(file_size, 1)) * 100))
                progress_state["input_rows_processed"] = stats["processed_rows"]
                progress_state["throughput_rps"] = rps
                progress_state["eta_sec"] = int(remaining_rows / max(rps, 1))
                
                push_log(log_queue, "INFO", f"Processed Chunk {chunk_idx} | Scanned: {stats['processed_rows']:,} | Retained: {stats['retained_rows']:,} | RPS: {rps:,}")
        except PermissionError:
            push_log(log_queue, "ERROR", f"Permission Denied. Please close {os.path.basename(input_file)} if it is open in Excel.")
            return False
            
        progress_state["final_total"] = stats["retained_rows"]
        push_log(log_queue, "SUCCESS", f"Cleaning completed successfully! File saved to {out_file}")
        return True

    except Exception as e:
        push_log(log_queue, "ERROR", f"Data Cleaning Engine Crashed: {str(e)}\n{traceback.format_exc()}")
        return False


# =========================================================
# NAME CLEANER & CATEGORIZER
# =========================================================
VALID_FOUND_VALUES = {"found"}

STRONG_ADDRESS_WORDS = {
    "floor", "flr", "road", "rd", "street", "st", "cross", "lane",
    "colony", "building", "bldg", "apartment", "apt", "flat", "house",
    "sector", "plot", "near", "behind", "opp", "opposite",
    "phase", "block", "area", "district", "village", "post",
    "taluk", "tehsil", "pincode", "pin", "door", "no", "doorno"
}

BUSINESS_WORDS = {
    "school", "college", "hospital", "clinic", "medical", "pharmacy",
    "store", "mart", "wear", "menswear", "fashion", "boutique",
    "cleaning", "solution", "enterprises", "enterprise", "agency",
    "agencies", "traders", "trading", "hardware", "software",
    "boarding", "international", "club", "restaurant", "hotel",
    "lodge", "saloon", "salon", "shop", "services", "service",
    "finance", "bank", "automobiles", "automobile", "motors",
    "stationery", "xerox", "mobiles", "mobile", "electronics",
    "jewellers", "jeweler", "jewellery", "bakery", "foods",
    "food", "cafe", "centre", "center", "academy"
}

GENERIC_NON_NAME_WORDS = {
    "guest", "customer", "unknown", "test", "demo", "sample",
    "null", "none", "na", "n/a", "dummy"
}

HONORIFICS = {
    "mr", "mrs", "ms", "miss", "dr", "prof", "sir", "madam", "shri", "smt"
}

def nc_normalize_text(value):
    if pd.isna(value): return ""
    text = str(value).replace("\ufeff", " ").strip()
    text = re.sub(r"\s+", " ", text)
    return text.strip()

def transliterate_to_english(text):
    if HAS_UNIDECODE: return unidecode(text).strip()
    return text.strip()

def clean_basic_name_text(text):
    text = re.sub(r"[^A-Za-z\s]", " ", text)
    return re.sub(r"\s+", " ", text).strip()

def title_case_name(text):
    return " ".join(w.capitalize() for w in text.split())

def token_stats(text):
    words = text.split()
    return words, len(words)

def count_digits(text):
    return len(re.findall(r"\d", text))

def contains_email_or_url(text):
    t = text.lower()
    if "@" in t: return True
    if "http://" in t or "https://" in t or "www." in t: return True
    return False

def count_word_hits(text, word_set):
    words = re.findall(r"[A-Za-z]+", text.lower())
    return sum(1 for w in words if w in word_set)

def dedupe_consecutive_and_global(words):
    compact = []
    for w in words:
        if not compact or compact[-1].lower() != w.lower():
            compact.append(w)
    seen = set()
    final = []
    for w in compact:
        lw = w.lower()
        if lw not in seen:
            final.append(w)
            seen.add(lw)
    return final

def strip_honorifics(words):
    return [w for w in words if w.lower() not in HONORIFICS]

def looks_like_repeated_phrase(original_text, cleaned_words):
    if len(cleaned_words) >= 4 and len(cleaned_words) % 2 == 0:
        half = len(cleaned_words) // 2
        if [w.lower() for w in cleaned_words[:half]] == [w.lower() for w in cleaned_words[half:]]:
            return True
    return False

def is_probable_business_or_address(original_text, cleaned_text):
    original_l = original_text.lower()
    cleaned_l = cleaned_text.lower()

    address_hits = count_word_hits(original_l, STRONG_ADDRESS_WORDS)
    business_hits = count_word_hits(original_l, BUSINESS_WORDS)
    digit_count = count_digits(original_text)
    comma_count = original_text.count(",")

    if business_hits >= 2: return True, "business_name"
    if address_hits >= 2: return True, "looks_like_address"
    if address_hits >= 1 and (digit_count >= 1 or comma_count >= 1): return True, "looks_like_address"
    if digit_count >= 4 and (address_hits >= 1 or business_hits >= 1): return True, "looks_like_address"
    if len(cleaned_text.split()) > 6 and (comma_count >= 1 or address_hits >= 1 or business_hits >= 1): return True, "too_long_non_name"
    return False, ""

def clean_target_name(raw_name):
    original = nc_normalize_text(raw_name)
    if not original: return "", "empty_original"
    if contains_email_or_url(original): return "", "contains_email_or_url"

    transliterated = transliterate_to_english(original)
    cleaned = clean_basic_name_text(transliterated)
    if not cleaned: return "", "empty_after_cleaning"

    words = cleaned.split()
    words = strip_honorifics(words)
    if not words: return "", "only_honorific"

    if words and words[0].lower() in GENERIC_NON_NAME_WORDS:
        words = words[1:]

    if not words: return "", "generic_non_name"

    if looks_like_repeated_phrase(original, words):
        half = len(words) // 2
        words = words[:half]

    words = dedupe_consecutive_and_global(words)
    cleaned = " ".join(words)
    cleaned = re.sub(r"\s+", " ", cleaned).strip()
    cleaned = title_case_name(cleaned)

    if not cleaned: return "", "empty_after_cleaning"

    words, word_count = token_stats(cleaned)
    if count_digits(original) >= 6: return "", "contains_many_digits"
    if len(cleaned) <= 1: return "", "too_short"
    if word_count == 1 and words[0].lower() in GENERIC_NON_NAME_WORDS: return "", "generic_non_name"
    if all(len(w) == 1 for w in words): return "", "initials_only"
    if word_count > 6: return "", "too_long"

    bad_flag, reason = is_probable_business_or_address(original, cleaned)
    if bad_flag: return "", reason

    return cleaned, ""

def run_name_cleaner(file_path: str, col_mob: str, col_name: str, col_status: str, out_dir: str, log_queue: List[Dict]) -> bool:
    if not HAS_UNIDECODE:
        push_log(log_queue, "ERROR", "Missing dependency: unidecode. Please run 'pip install unidecode'")
        return False

    try:
        ensure_dir(out_dir)
        push_log(log_queue, "INFO", f"Reading Data File: {file_path}")
        
        base_name = os.path.splitext(os.path.basename(file_path))[0]
        
        c_file = os.path.join(out_dir, f"{base_name}_consolidate.csv")
        cn_file = os.path.join(out_dir, f"{base_name}_clean_with_not_found.csv")
        o_file = os.path.join(out_dir, f"{base_name}_only_clean.csv")
        n_file = os.path.join(out_dir, f"{base_name}_not_found_data.csv")
        s_file = os.path.join(out_dir, f"{base_name}_summary.json")

        c_writer = CsvChunkWriter(c_file, ["mobile", "lookup_response", "cleaned_name", "status", "status_bucket", "is_clean", "removed_reason"])
        cn_writer = CsvChunkWriter(cn_file, ["mobile", "name", "lookup_response", "status", "status_bucket"])
        o_writer = CsvChunkWriter(o_file, ["mobile", "name", "lookup_response", "status"])
        n_writer = CsvChunkWriter(n_file, ["mobile", "lookup_response", "cleaned_name", "status", "is_clean", "removed_reason"])

        totals = {"total_rows": 0, "clean_names": 0, "non_clean_names": 0}
        status_counter = {"Found": 0, "Not Found": 0}
        removal_reason_counter = {}
        file_counts = {"consolidate": 0, "clean_with_not_found": 0, "only_clean": 0, "not_found_data": 0}
        
        start_time = time.time()
        
        try:
            for chunk_idx, chunk in enumerate(iter_file_chunks(file_path, INPUT_CHUNK_SIZE), start=1):
                if not {col_mob, col_name, col_status}.issubset(set(chunk.columns)):
                    push_log(log_queue, "ERROR", f"Missing mapped columns in chunk {chunk_idx}")
                    return False

                c_rows, cn_rows, o_rows, n_rows = [], [], [], []

                for _, row in chunk.iterrows():
                    totals["total_rows"] += 1
                    mobile = nc_normalize_text(row.get(col_mob, ""))
                    lookup_response = nc_normalize_text(row.get(col_name, ""))
                    raw_status = nc_normalize_text(row.get(col_status, ""))
                    found_flag = raw_status.lower() in VALID_FOUND_VALUES

                    status_bucket = "Found" if found_flag else "Not Found"
                    status_counter[status_bucket] += 1

                    cleaned_name, removal_reason = clean_target_name(lookup_response)
                    is_clean = bool(cleaned_name)

                    if is_clean:
                        totals["clean_names"] += 1
                    else:
                        totals["non_clean_names"] += 1
                        removal_reason_counter[removal_reason] = removal_reason_counter.get(removal_reason, 0) + 1

                    c_rows.append([mobile, lookup_response, cleaned_name, raw_status, status_bucket, "Yes" if is_clean else "No", "" if is_clean else removal_reason])
                    file_counts["consolidate"] += 1

                    if is_clean:
                        cn_rows.append([mobile, cleaned_name, lookup_response, raw_status, status_bucket])
                        file_counts["clean_with_not_found"] += 1

                    if is_clean and found_flag:
                        o_rows.append([mobile, cleaned_name, lookup_response, raw_status])
                        file_counts["only_clean"] += 1

                    if not found_flag:
                        n_rows.append([mobile, lookup_response, cleaned_name, raw_status, "Yes" if is_clean else "No", "" if is_clean else removal_reason])
                        file_counts["not_found_data"] += 1

                c_writer.write_df(pd.DataFrame(c_rows, columns=c_writer.headers))
                cn_writer.write_df(pd.DataFrame(cn_rows, columns=cn_writer.headers))
                o_writer.write_df(pd.DataFrame(o_rows, columns=o_writer.headers))
                n_writer.write_df(pd.DataFrame(n_rows, columns=n_writer.headers))

                rps = int(totals['total_rows'] / max(time.time() - start_time, 1))
                push_log(log_queue, "INFO", f"Processed Chunk {chunk_idx} | Rows: {totals['total_rows']:,} | RPS: {rps:,}")
        except PermissionError:
            push_log(log_queue, "ERROR", f"Permission Denied. Please close {os.path.basename(file_path)} if it is open in Excel.")
            return False

        summary = {
            "input_file": file_path,
            "output_files": {"consolidate": c_file, "clean_with_not_found": cn_file, "only_clean": o_file, "not_found_data": n_file},
            "overall_summary": {
                "total_rows": totals["total_rows"],
                "found_rows": status_counter["Found"],
                "not_found_rows": status_counter["Not Found"],
                "clean_names_total": totals["clean_names"],
                "non_clean_names_total": totals["non_clean_names"]
            },
            "file_wise_summary": {
                "consolidate": {"rows": file_counts["consolidate"]},
                "clean_with_not_found": {"rows": file_counts["clean_with_not_found"]},
                "only_clean": {"rows": file_counts["only_clean"]},
                "not_found_data": {"rows": file_counts["not_found_data"]}
            },
            "non_clean_reason_breakdown": removal_reason_counter,
            "status_breakdown": status_counter
        }

        with open(s_file, "w", encoding="utf-8") as f:
            json.dump(summary, f, indent=4, ensure_ascii=False)

        push_log(log_queue, "SUCCESS", f"Processing completed successfully! Outputs saved to {out_dir}")
        return True

    except Exception as e:
        push_log(log_queue, "ERROR", f"Name Cleaning Job Crashed: {str(e)}\n{traceback.format_exc()}")
        return False

# =========================================================
# BASE64 EXTRACTOR
# =========================================================
def run_base64_extractor(master_folder: str, output_folder: str, log_queue: List[Dict], progress_state: Dict[str, Any]) -> bool:
    try:
        push_log(log_queue, "INFO", f"Starting Base64 Extraction from: {master_folder}")
        if not os.path.exists(master_folder):
            push_log(log_queue, "ERROR", "Master folder not found.")
            return False
            
        ensure_dir(output_folder)
        
        files_to_process = []
        for root, dirs, files in os.walk(master_folder):
            for filename in files:
                if filename.endswith('.txt'):
                    files_to_process.append(os.path.join(root, filename))
                    
        total_files = len(files_to_process)
        if total_files == 0:
            push_log(log_queue, "WARN", "No .txt files found to process.")
            return True
            
        push_log(log_queue, "INFO", f"Found {total_files} .txt files. Processing...")
        
        success_count = 0
        valid_chars = set('ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz0123456789+/=')
        
        start_time = time.time()
        
        for idx, file_path in enumerate(files_to_process, 1):
            folder_name = os.path.basename(os.path.dirname(file_path))
            filename = os.path.basename(file_path)
            
            try:
                with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                    content = f.read()
                    
                content = content.replace('<img src="', '') 
                content = content.replace('data:image/jpeg;base64,', '')
                content = content.replace('data:image/png;base64,', '')
                content = content.replace('data:application/pdf;base64,', '')
                
                clean_b64 = "".join(c for c in content if c in valid_chars)
                
                if not clean_b64:
                    continue
                    
                if clean_b64.startswith('JVBERi0'): ext = '.pdf'
                elif clean_b64.startswith('/9j/'): ext = '.jpg'
                elif clean_b64.startswith('iVBORw0KGgo'): ext = '.png'
                else: ext = '.bin'
                
                missing_padding = len(clean_b64) % 4
                if missing_padding != 0:
                    clean_b64 += '=' * (4 - missing_padding)
                    
                base_name = folder_name
                out_file_path = os.path.join(output_folder, base_name + ext)
                counter = 1
                while os.path.exists(out_file_path):
                    out_file_path = os.path.join(output_folder, f"{base_name}_{counter}{ext}")
                    counter += 1
                    
                file_data = base64.b64decode(clean_b64, validate=True)
                with open(out_file_path, 'wb') as out_file:
                    out_file.write(file_data)
                    
                success_count += 1
                if idx % 50 == 0:
                    rps = int(idx / max(time.time() - start_time, 1))
                    progress_state["input_rows_processed"] = idx
                    progress_state["throughput_rps"] = rps
                    push_log(log_queue, "INFO", f"Processed {idx}/{total_files} files...")
                    
            except Exception as e:
                push_log(log_queue, "WARN", f"Error decoding {filename}: {str(e)}")
                
        push_log(log_queue, "SUCCESS", f"Extraction complete! Successfully decoded {success_count} files to {output_folder}")
        return True
        
    except Exception as e:
        push_log(log_queue, "ERROR", f"Base64 Extractor Crashed: {str(e)}\n{traceback.format_exc()}")
        return False

# =========================================================
# EDUCATION DOCUMENTS OCR
# =========================================================

def ocr_clean_text(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()

def ocr_safe_str(value: Any) -> str:
    if value is None: return ""
    return str(value).strip()

def pil_to_cv_bgr(img: Image.Image) -> np.ndarray:
    arr = np.array(img.convert("RGB"))
    return cv2.cvtColor(arr, cv2.COLOR_RGB2BGR)

def render_pdf_to_images(pdf_path: str, scale: float = 1.5) -> List[Image.Image]:
    if not HAS_PDFIUM:
        raise ImportError("pypdfium2 is required for PDF support. Install it with: pip install pypdfium2")
    images = []
    pdf = pdfium.PdfDocument(pdf_path)
    try:
        for i in range(len(pdf)):
            page = pdf[i]
            bitmap = page.render(scale=scale)
            pil_image = bitmap.to_pil().convert("RGB")
            images.append(pil_image)
    finally:
        pdf.close()
    return images

def load_images_from_file(file_path: str) -> List[Image.Image]:
    ext = os.path.splitext(file_path.lower())[1]
    if ext == ".pdf":
        return render_pdf_to_images(file_path)
    return [Image.open(file_path).convert("RGB")]

def downscale_image(img: Image.Image, max_side: int = 1600) -> Image.Image:
    w, h = img.size
    longest = max(w, h)
    if longest <= max_side: return img
    scale = max_side / float(longest)
    new_w = max(1, int(w * scale))
    new_h = max(1, int(h * scale))
    return img.resize((new_w, new_h), Image.LANCZOS)

def preprocess_image_for_ocr(img: Image.Image) -> Image.Image:
    if not HAS_CV2: return img
    rgb = np.array(img.convert("RGB"))
    gray = cv2.cvtColor(rgb, cv2.COLOR_RGB2GRAY)
    gray = cv2.GaussianBlur(gray, (3, 3), 0)
    processed = cv2.adaptiveThreshold(gray, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C, cv2.THRESH_BINARY, 31, 15)
    kernel = np.ones((1, 1), np.uint8)
    processed = cv2.morphologyEx(processed, cv2.MORPH_OPEN, kernel)
    return Image.fromarray(processed).convert("RGB")

def extract_doctr_text_and_confidence(result) -> Tuple[str, float, List[Dict[str, Any]]]:
    all_lines = []
    confidences = []
    raw_records = []
    for page_idx, page in enumerate(result.pages, start=1):
        for block in page.blocks:
            for line in block.lines:
                words = []
                for word in line.words:
                    value = ocr_safe_str(word.value)
                    conf = float(word.confidence) if word.confidence is not None else 0.0
                    if value:
                        words.append(value)
                        confidences.append(conf)
                        raw_records.append({
                            "page": page_idx, "text": value, "confidence": round(conf, 4), "bbox": word.geometry
                        })
                line_text = " ".join(words).strip()
                if line_text: all_lines.append(line_text)

    full_text = ocr_clean_text("\n".join(all_lines))
    avg_confidence = round(sum(confidences) / len(confidences), 4) if confidences else 0.0
    return full_text, avg_confidence, raw_records

def run_doctr_on_images(predictor, images: List[Image.Image]) -> Tuple[str, float, List[Dict[str, Any]]]:
    doctr_inputs = [pil_to_cv_bgr(img) for img in images]
    if HAS_TORCH:
        with torch.inference_mode():
            result = predictor(doctr_inputs)
    else:
        result = predictor(doctr_inputs)
    return extract_doctr_text_and_confidence(result)

def build_extraction_prompt(ocr_text: str) -> str:
    return f"""
You are an expert document information extraction engine.
Your task is to accurately extract educational details from the provided OCR text of a certificate or marksheet.

Extract exactly these keys in strict JSON format:
{{
  "Roll No": "",
  "University Name": "",
  "Course Name": "",
  "Passing Year": "",
  "Candidate Name": "",
  "Passing Status": "",
  "Institute Name": "",
  "Confidence": 0.0
}}

Rules:
1. Use ONLY the OCR text. Do not invent or guess values.
2. If a field is truly missing, return an empty string "".
3. Passing Year MUST be the 4-digit year the student took the exam or graduated (e.g., 2014, 2018, 2024). YOU MUST IGNORE years associated with Acts, Laws, or establishment dates (e.g., 1956, 1961, 2005, 1949).
4. Passing Status must be exactly one of: Pass, Fail, Completed, Promoted, Appeared, Unknown, Needs Improvement.
5. Candidate Name should be the student's name. YOU MUST IGNORE parent names (usually preceded by Father, Mother, Shri, Smt, or Son/Daughter of).
6. Institute Name should be the specific school, college, or study centre they attended. 
7. Roll No should be the student's specific Seat No, Roll No, Registration No, PRN, or Enrolment No.
8. Course Name is the Degree, Trade, or Examination name (e.g., Bachelor of Arts, Master of Business Administration, Senior Secondary Examination, Electrician).

OCR TEXT:
\"\"\"
{ocr_text}
\"\"\"
""".strip()

def ollama_extract_fields(ocr_text: str, model_name: str, timeout: int = 180, n_predict: int = 256) -> Dict[str, Any]:
    prompt = build_extraction_prompt(ocr_text[:6000])
    payload = {
        "model": model_name,
        "prompt": prompt,
        "stream": False,
        "format": "json",
        "options": {"temperature": 0.0, "num_predict": n_predict}
    }
    try:
        response = requests.post("http://localhost:11434/api/generate", json=payload, timeout=timeout)
        response.raise_for_status()
        result_json = response.json()
        raw_response = result_json.get("response", "{}")
        parsed = json.loads(raw_response)
        
        out = {
            "Roll No": ocr_safe_str(parsed.get("Roll No", "")),
            "University Name": ocr_safe_str(parsed.get("University Name", "")),
            "Course Name": ocr_safe_str(parsed.get("Course Name", "")),
            "Passing Year": ocr_safe_str(parsed.get("Passing Year", "")),
            "Candidate Name": ocr_safe_str(parsed.get("Candidate Name", "")),
            "Passing Status": ocr_safe_str(parsed.get("Passing Status", "Unknown")),
            "Institute Name": ocr_safe_str(parsed.get("Institute Name", "")),
            "Confidence": parsed.get("Confidence", 0.0),
        }
        return out
    except requests.exceptions.RequestException as e:
        raise RuntimeError(f"Ollama API request failed: {str(e)}")
    except json.JSONDecodeError:
        raise ValueError(f"Ollama returned invalid JSON.")

def is_education_document(ocr_text: str) -> bool:
    text = ocr_text.lower()
    keywords = [
        "university", "board", "council", "statement of marks", "marksheet",
        "admit card", "hall ticket", "roll", "registration", "seat number",
        "prn", "examination", "result", "percentage", "student", "course",
        "semester", "grade", "school", "college", "institute", "trade",
        "discipline", "apprenticeship", "degree", "diploma", "bachelor", "master",
        "ncvt", "iti", "industrial training", "b.com", "bachelor of commerce"
    ]
    score = sum(1 for k in keywords if k in text)
    return score >= 2

def ocr_normalize_line(line: str) -> str:
    line = ocr_clean_text(line)
    return re.sub(r"[|]+", " ", line).strip()

def get_clean_lines(ocr_text: str) -> List[str]:
    lines = [ocr_normalize_line(line) for line in ocr_text.splitlines()]
    return [line for line in lines if line and not line.startswith("--- Page")]

def cleanup_extracted_value(value: str) -> str:
    value = ocr_safe_str(value)
    value = re.sub(r"\s{2,}", " ", value)
    return value.strip(" :-;,._")

def looks_like_name(value: str) -> bool:
    value = cleanup_extracted_value(value)
    if not value or len(value.split()) < 2: return False
    if any(ch.isdigit() for ch in value): return False
    bad = ["university", "board", "council", "school", "college", "institute", "course", "result", "examination", "admit"]
    if any(b in value.lower() for b in bad): return False
    return True

def looks_like_roll(value: str) -> bool:
    value = cleanup_extracted_value(value)
    if not value or len(value) < 4: return False
    if not re.search(r"[A-Z0-9]", value, re.IGNORECASE): return False
    if value.lower() in {"name", "centre", "center", "course", "code", "index", "number"}: return False
    return True

def looks_like_institute(value: str) -> bool:
    return any(k in value.lower() for k in ["school", "college", "institute", "center", "centre", "vidyapith", "academy", "secondary", "polytechnic", "campus"])

def first_regex_match(lines: List[str], patterns: List[str], exclude_keywords: List[str] = None) -> str:
    exclude_keywords = exclude_keywords or []
    for line in lines:
        if any(ex.lower() in line.lower() for ex in exclude_keywords): continue
        for pattern in patterns:
            m = re.search(pattern, line, flags=re.IGNORECASE)
            if m:
                value = cleanup_extracted_value(m.group(1))
                if value: return value
    return ""

def detect_roll_no(lines: List[str]) -> str:
    patterns = [
        r"seat\s*number\s*[:\-]?\s*([A-Z0-9\/\-\s]+)", r"seat\s*no\.?\s*[:\-]?\s*([A-Z0-9\/\-\s]+)",
        r"roll\s*no\.?\s*[:\-]?\s*([A-Z0-9\/\-\s]+)", r"roll\s*[:\-]?\s*([A-Z0-9\/\-\s]+)",
        r"registration\s*no\.?\s*[:\-]?\s*([A-Z0-9\/\-\s]+)", r"enrollment\s*no\.?\s*[:\-]?\s*([A-Z0-9\/\-\s]+)",
        r"s\.i\.d\.\s*no\.?\s*[:\-]?\s*([A-Z0-9\/\-\s]+)",
    ]
    value = first_regex_match(lines, patterns)
    if value and looks_like_roll(value): return value.strip()
    for idx, line in enumerate(lines):
        low = line.lower().strip()
        if low in ["seat no.", "seat no", "roll", "roll no", "roll no.", "s.i.d. no."]:
            if idx + 1 < len(lines):
                nxt = cleanup_extracted_value(lines[idx + 1])
                if looks_like_roll(nxt): return nxt
    return ""

def detect_candidate_name(lines: List[str]) -> str:
    patterns = [
        r"candidate'?s?\s*name\s*[:\-]?\s*(.+)", r"student'?s?\s*name\s*[:\-]?\s*(.+)",
        r"name\s+of\s+the\s+candidate\s*[:\-]?\s*(.+)", r"shri\s*/\s*smt\.?\s*[:\-]?\s*(.+)", r"^name\s*[:\-]?\s*(.+)",
    ]
    value = first_regex_match(lines, patterns, exclude_keywords=["mother", "father", "guardian", "course", "institute", "university"])
    if looks_like_name(value): return value
    for line in lines:
        candidate = cleanup_extracted_value(line)
        if looks_like_name(candidate):
            alpha_count = max(1, sum(1 for ch in candidate if ch.isalpha()))
            upper_ratio = sum(1 for ch in candidate if ch.isupper()) / alpha_count
            if upper_ratio > 0.8 and len(candidate.split()) <= 4:
                if not any(bad in candidate.lower() for bad in ["university", "board", "college", "education", "distance"]):
                    return candidate
    return ""

def detect_university_name(lines: List[str]) -> str:
    candidates = []
    for line in lines:
        low = line.lower()
        if any(k in low for k in ["university", "board", "council", "open university"]):
            cleaned = cleanup_extracted_value(line)
            word_count = len(cleaned.split())
            if 2 <= word_count <= 10 and not any(bad in low for bad in ["admitted by", "members of", "recognised by", "chancellor"]):
                candidates.append(cleaned)
    if not candidates: return ""
    return sorted(candidates, key=lambda x: (-len(x), x))[0]

def detect_course_name(lines: List[str]) -> str:
    patterns = [
        r"statement of marks for (.+?) examination", r"discipline\s*[:\-]?\s*(.+)",
        r"stream\s*[:\-]?\s*(.+)", r"course\s*name\s*[:\-]?\s*(.+)",
        r"programme\s*[:\-]?\s*(.+)", r"degree\s*[:\-]?\s*(.+)", r"trade\s*[:\-]?\s*(.+)",
    ]
    value = first_regex_match(lines, patterns)
    value = cleanup_extracted_value(value)
    if value and value.lower() not in {"name", "result", "pass", "fail", "centre", "center", "(theoretical)"}:
        return value
    for idx, line in enumerate(lines):
        low = line.lower().strip()
        if low in ["discipline", "stream", "course", "trade", "programme"]:
            if idx + 1 < len(lines):
                nxt = cleanup_extracted_value(lines[idx + 1])
                if len(nxt) > 2: return nxt
    return ""

def detect_institute_name(lines: List[str]) -> str:
    patterns = [
        r"study\s*center\s*[:\-]?\s*(.+)", r"study\s*centre\s*[:\-]?\s*(.+)",
        r"exam\s*center\s*[:\-]?\s*(.+)", r"exam\s*centre\s*[:\-]?\s*(.+)",
        r"examination\s*centre\s*[:\-]?\s*(.+)", r"center\s*[:;\-]?\s*(.+)",
        r"centre\s*[:;\-]?\s*(.+)", r"institute\s*[:\-]?\s*(.+)", r"college\s*[:\-]?\s*(.+)", r"school\s*[:\-]?\s*(.+)",
    ]
    value = first_regex_match(lines, patterns)
    value = cleanup_extracted_value(value)
    if value.lower() in ["(theoretical)", "practical", "code"]: value = ""
    if looks_like_institute(value): return value
    for idx, line in enumerate(lines):
        low = line.lower()
        if any(keyword in low for keyword in ["centre", "center", "school", "college", "institution"]):
            for offset in [1, 2]:
                if idx + offset < len(lines):
                    nxt = cleanup_extracted_value(lines[idx + offset])
                    nxt_alpha = re.sub(r"^\d+\s*", "", nxt)
                    if looks_like_institute(nxt_alpha): return nxt_alpha
    return ""

def detect_passing_status(lines: List[str]) -> str:
    blob = " ".join(lines).lower()
    if "needs improvement" in blob: return "Needs Improvement"
    if "result: pass" in blob or re.search(r"\bpass\b", blob) or re.search(r"\bqualified\b", blob): return "Pass"
    if "result: fail" in blob or re.search(r"\bfail\b", blob): return "Fail"
    if re.search(r"\bpromoted\b", blob): return "Promoted"
    if re.search(r"\bappeared\b", blob): return "Appeared"
    if re.search(r"\bcompleted\b", blob): return "Completed"
    return "Unknown"

def detect_passing_year(lines: List[str]) -> str:
    priority_keywords = ["year of passing", "passing year", "date:", "dated", "examination", "statement of marks", "admit card", "result", "issued", "semester"]
    priority_lines = [line for line in lines if any(k in line.lower() for k in priority_keywords)]
    year_pattern = r"\b(19\d{2}|20\d{2}|21\d{2})\b"
    for line in priority_lines + lines:
        m = re.search(year_pattern, line)
        if m: return m.group(1)
    return ""

def smart_extract_fields(ocr_text: str) -> Dict[str, Any]:
    lines = get_clean_lines(ocr_text)
    roll_no = detect_roll_no(lines)
    university_name = detect_university_name(lines)
    course_name = detect_course_name(lines)
    passing_year = detect_passing_year(lines)
    candidate_name = detect_candidate_name(lines)
    passing_status = detect_passing_status(lines)
    institute_name = detect_institute_name(lines)

    fields = [roll_no, university_name, course_name, passing_year, candidate_name, institute_name]
    filled_count = sum(1 for f in fields if f)
    dynamic_conf = round(filled_count / max(len(fields), 1), 2)

    return {
        "Roll No": roll_no, "University Name": university_name, "Course Name": course_name,
        "Passing Year": passing_year, "Candidate Name": candidate_name, "Passing Status": passing_status,
        "Institute Name": institute_name, "Confidence": dynamic_conf,
    }

def combine_confidence(ocr_conf: float, extract_conf: float) -> float:
    try: extract_conf = float(extract_conf)
    except Exception: extract_conf = 0.0
    final_conf = (0.7 * ocr_conf) + (0.3 * extract_conf)
    return round(max(0.0, min(1.0, final_conf)), 4)

def needs_review(row: Dict[str, Any]) -> bool:
    required = ["Candidate Name", "Passing Year"]
    if any(not ocr_safe_str(row.get(k)) for k in required): return True
    if float(row.get("Confidence", 0.0)) < 0.60: return True
    if "Skipped" in ocr_safe_str(row.get("Status")): return True
    if "error" in ocr_safe_str(row.get("Extraction Mode")): return True
    return False

def should_use_llm(extracted_data: Dict[str, Any], ocr_confidence: float, ocr_text: str, enable_llm: bool = False) -> bool:
    if enable_llm and len(ocr_text) >= 80: return True
    return False

def normalize_final_row(row: Dict[str, Any]) -> Dict[str, Any]:
    out = dict(row)
    year_match = re.search(r"\b(19\d{2}|20\d{2}|21\d{2})\b", ocr_safe_str(out.get("Passing Year", "")))
    out["Passing Year"] = year_match.group(1) if year_match else ""

    valid_statuses = {"Pass", "Fail", "Completed", "Promoted", "Appeared", "Unknown", "Needs Improvement"}
    status = ocr_safe_str(out.get("Passing Status", "Unknown"))
    if status not in valid_statuses:
        low = status.lower()
        if "pass" in low or "qualified" in low: status = "Pass"
        elif "fail" in low: status = "Fail"
        elif "complete" in low: status = "Completed"
        elif "promot" in low: status = "Promoted"
        elif "appear" in low: status = "Appeared"
        elif "needs improvement" in low: status = "Needs Improvement"
        else: status = "Unknown"
    out["Passing Status"] = status

    for key in ["Roll No", "University Name", "Course Name", "Candidate Name", "Institute Name"]:
        out[key] = cleanup_extracted_value(out.get(key, ""))
    return out

def ocr_process_file(predictor, file_path: str, ollama_model: str, preprocess: bool, max_side: int, keep_raw_ocr: bool, enable_llm: bool) -> Dict[str, Any]:
    file_name = os.path.basename(file_path)
    try:
        stage_times = {}
        t0 = time.time()
        images = load_images_from_file(file_path)
        stage_times["load_sec"] = round(time.time() - t0, 2)

        t1 = time.time()
        images = [downscale_image(img, max_side=max_side) for img in images]
        stage_times["resize_sec"] = round(time.time() - t1, 2)

        if preprocess:
            t2 = time.time()
            images = [preprocess_image_for_ocr(img) for img in images]
            stage_times["preprocess_sec"] = round(time.time() - t2, 2)
        else:
            stage_times["preprocess_sec"] = 0.0

        t3 = time.time()
        ocr_text, ocr_confidence, raw_ocr = run_doctr_on_images(predictor, images)
        stage_times["ocr_sec"] = round(time.time() - t3, 2)

        row = {
            "File Name": file_name, "Roll No": "", "University Name": "", "Course Name": "", "Passing Year": "", 
            "Candidate Name": "", "Passing Status": "Unknown", "Institute Name": "", "Confidence": 0.0, "OCR Confidence": round(ocr_confidence, 4),
            "OCR Text": ocr_text if keep_raw_ocr else "", "Status": "Skipped: not an education document", "Extraction Mode": "skipped",
            "Review Required": True, "Load Sec": stage_times["load_sec"], "Resize Sec": stage_times["resize_sec"], "Preprocess Sec": stage_times["preprocess_sec"],
            "OCR Sec": stage_times["ocr_sec"], "LLM Sec": 0.0, "Raw OCR": raw_ocr if keep_raw_ocr else [],
        }

        if not ocr_text:
            row["Status"] = "No text extracted"
            row["Extraction Mode"] = "none"
            return row

        ocr_text = ocr_text[:12000]
        if not is_education_document(ocr_text): return row

        extracted = smart_extract_fields(ocr_text)
        extraction_mode = "smart_rules"
        llm_time = 0.0
        status = "Success"

        if should_use_llm(extracted, ocr_confidence, ocr_text, enable_llm=enable_llm):
            try:
                t4 = time.time()
                llm_data = ollama_extract_fields(ocr_text=ocr_text, model_name=ollama_model)
                llm_time = round(time.time() - t4, 2)
                extracted = llm_data
                extraction_mode = "ollama"
            except Exception as llm_error:
                status = f"Processed with smart rules (Ollama API issue: {str(llm_error)})"
                extraction_mode = "ollama_error"

        extracted = normalize_final_row(extracted)
        for k in ["Roll No", "University Name", "Course Name", "Passing Year", "Candidate Name", "Passing Status", "Institute Name"]:
            row[k] = ocr_safe_str(extracted.get(k, ""))
        
        row["Confidence"] = combine_confidence(ocr_confidence, extracted.get("Confidence", 0.0))
        row["Status"] = status
        row["Extraction Mode"] = extraction_mode
        row["LLM Sec"] = llm_time
        row["Review Required"] = needs_review(row)

        return row
    except Exception as e:
        return {
            "File Name": file_name, "Roll No": "", "University Name": "", "Course Name": "", "Passing Year": "", 
            "Candidate Name": "", "Passing Status": "Unknown", "Institute Name": "", "Confidence": 0.0, "OCR Confidence": 0.0, "OCR Text": "",
            "Status": f"Error: {str(e)}", "Extraction Mode": "error", "Review Required": True, "Load Sec": 0.0, "Resize Sec": 0.0, "Preprocess Sec": 0.0, "OCR Sec": 0.0, "LLM Sec": 0.0, "Raw OCR": [],
        }

def run_education_ocr(folder_path: str, enable_llm: bool, ollama_model: str, preprocess: bool, max_side: int, keep_raw_ocr: bool, log_queue: List[Dict], progress_state: Dict[str, Any]) -> bool:
    if not HAS_DOCTR or not HAS_CV2:
        push_log(log_queue, "ERROR", "Missing required libraries. Please install: pip install python-doctr opencv-python pypdfium2 torch torchvision")
        return False
    
    if enable_llm:
        try:
            resp = requests.get("http://localhost:11434/")
            resp.raise_for_status()
        except requests.exceptions.RequestException:
            push_log(log_queue, "ERROR", "Ollama is not running! Please open the Ollama app or start the service.")
            return False

    if HAS_TORCH:
        try:
            torch.set_num_threads(4)
            torch.set_grad_enabled(False)
        except Exception: pass

    try:
        ensure_dir(folder_path)
        files = []
        for name in os.listdir(folder_path):
            full_path = os.path.join(folder_path, name)
            if os.path.isfile(full_path) and os.path.splitext(name.lower())[1] in SUPPORTED_EXTENSIONS:
                files.append(full_path)
        
        if not files:
            push_log(log_queue, "WARN", "No supported documents found in the folder.")
            return True

        push_log(log_queue, "INFO", f"Loading docTR OCR model for {len(files)} files...")
        predictor = ocr_predictor(det_arch="fast_base", reco_arch="crnn_vgg16_bn", pretrained=True)
        push_log(log_queue, "SUCCESS", "docTR model loaded successfully.")

        csv_out = os.path.join(folder_path, "education_output.csv")
        review_out = os.path.join(folder_path, "education_review.csv")
        
        fields = ["File Name", "Roll No", "University Name", "Course Name", "Passing Year", "Candidate Name", "Passing Status", "Institute Name", "Confidence", "OCR Confidence", "Status", "Extraction Mode", "Review Required", "Load Sec", "Resize Sec", "Preprocess Sec", "OCR Sec", "LLM Sec"]
        
        with open(csv_out, "w", newline="", encoding="utf-8") as f:
            csv.DictWriter(f, fieldnames=fields).writeheader()
        with open(review_out, "w", newline="", encoding="utf-8") as f:
            csv.DictWriter(f, fieldnames=fields).writeheader()

        total_files = len(files)
        processed_times = []
        review_count = 0
        
        for idx, file_path in enumerate(files, start=1):
            file_name = os.path.basename(file_path)
            
            file_start = time.time()
            result = ocr_process_file(predictor, file_path, ollama_model, preprocess, max_side, keep_raw_ocr, enable_llm)
            file_elapsed = round(time.time() - file_start, 2)
            processed_times.append(file_elapsed)
            
            with open(csv_out, "a", newline="", encoding="utf-8") as f:
                csv.DictWriter(f, fieldnames=fields).writerow({k: result.get(k, "") for k in fields})
            
            if result.get("Review Required", False):
                review_count += 1
                with open(review_out, "a", newline="", encoding="utf-8") as f:
                    csv.DictWriter(f, fieldnames=fields).writerow({k: result.get(k, "") for k in fields})

            avg_time = sum(processed_times) / len(processed_times)
            eta = avg_time * (total_files - idx)
            
            progress_state["input_rows_processed"] = idx
            progress_state["input_percent"] = int(idx / total_files * 100)
            progress_state["throughput_rps"] = round(avg_time, 2)
            progress_state["eta_sec"] = int(eta)
            progress_state["review_count"] = review_count

            push_log(log_queue, "INFO", f"[{idx}/{total_files}] Processed {file_name} in {file_elapsed}s | Review: {result.get('Review Required')}")
        
        push_log(log_queue, "SUCCESS", f"OCR Process completed. Main Output: {csv_out}")
        return True

    except Exception as e:
        push_log(log_queue, "ERROR", f"Education OCR Engine Crashed: {str(e)}\n{traceback.format_exc()}")
        return False